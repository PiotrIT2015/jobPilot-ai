import io
import re
from pathlib import Path

from docx import Document
from pypdf import PdfReader


MAX_TEXT_CHARS = 200_000

# Transparent starter list; this is keyword extraction, not an AI judgment.
SKILL_KEYWORDS = [
    "Python", "Java", "JavaScript", "TypeScript", "PHP", "SQL", "MySQL",
    "PostgreSQL", "Windows", "Linux", "Docker", "Kubernetes", "Git",
    "CI/CD", "Jenkins", "GitHub Actions", "FastAPI", "React", "AWS",
    "Azure", "TCP/IP", "DNS", "DHCP", "Active Directory", "Bash",
    "PowerShell", "LAMP", "Nmap", "CCNA", "IT Support", "DevOps",
]


def extract_text(data: bytes, extension: str) -> str:
    extension = extension.lower()
    if extension == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        chunks = [(page.extract_text() or "") for page in reader.pages]
        text = "\n".join(chunks)
    elif extension == ".docx":
        document = Document(io.BytesIO(data))
        chunks = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                chunks.append(" | ".join(cell.text.strip() for cell in row.cells))
        text = "\n".join(chunks)
    else:
        raise ValueError("Obsługiwane formaty to PDF i DOCX.")

    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        raise ValueError(
            "Nie udało się odczytać tekstu. Jeśli PDF jest skanem, potrzebny będzie OCR."
        )
    return text[:MAX_TEXT_CHARS]


def build_candidate_profile(text: str) -> dict:
    lower = text.casefold()
    found_skills = [skill for skill in SKILL_KEYWORDS if skill.casefold() in lower]

    # Keep this deliberately simple and explainable; later stage can add AI extraction.
    lines = [line.strip(" •\t-") for line in text.splitlines() if line.strip()]
    education_terms = ("education", "wykształcenie", "university", "uczelnia", "degree", "inż.")
    experience_terms = ("experience", "doświadczenie", "employment", "work history", "staż")

    education_lines = [line for line in lines if any(term in line.casefold() for term in education_terms)][:8]
    experience_lines = [line for line in lines if any(term in line.casefold() for term in experience_terms)][:8]

    return {
        "skills_detected": found_skills,
        "education_hints": education_lines,
        "experience_hints": experience_lines,
        "character_count": len(text),
        "note": "Wynik jest wstępną ekstrakcją słów kluczowych, a nie pełną oceną CV przez AI."
    }
